import os
import io
import re
import pandas as pd
from pypdf import PdfReader
try:
    import docx
except ImportError:
    docx = None

try:
    from PIL import Image
    import pytesseract
    HAS_PYTESSERACT = True
except ImportError:
    HAS_PYTESSERACT = False

# Maximum allowed CV upload size (15 MB)
MAX_FILE_SIZE_BYTES = 15 * 1024 * 1024
ALLOWED_EXTENSIONS = {".pdf", ".docx", ".doc", ".txt", ".csv"}

class CVExtractionError(Exception):
    """Custom exception raised when CV processing encounters a validation or parsing failure."""
    pass

def validate_cv_file(file_obj) -> None:
    """
    Validates file extension and size constraints to prevent processing oversized or invalid files.
    """
    filename = getattr(file_obj, "name", "").lower()
    ext = os.path.splitext(filename)[1]
    if ext not in ALLOWED_EXTENSIONS:
        raise CVExtractionError(
            f"Unsupported file format '{ext}'. Please upload a PDF (.pdf), Word document (.docx), CSV (.csv), or Text file (.txt)."
        )

    # Check file size if available
    size = getattr(file_obj, "size", None)
    if size is not None and size > MAX_FILE_SIZE_BYTES:
        raise CVExtractionError(
            f"File size exceeds the 15 MB limit ({round(size / (1024 * 1024), 1)} MB). Please upload a smaller file."
        )

def _attempt_ocr_on_pdf(reader: PdfReader, max_pages: int = 5) -> str:
    """
    Attempts optical character recognition (OCR) on scanned PDF pages using pytesseract.
    Fails gracefully if OCR binaries are missing or if pytesseract is not installed.
    """
    if not HAS_PYTESSERACT:
        return ""

    ocr_text = []
    try:
        pages_to_process = reader.pages[:max_pages]
        for page in pages_to_process:
            for image_file_object in page.images:
                try:
                    img = Image.open(io.BytesIO(image_file_object.data))
                    txt = pytesseract.image_to_string(img)
                    if txt.strip():
                        ocr_text.append(txt.strip())
                except Exception:
                    continue
    except Exception:
        # Tesseract binary not installed on host OS or other OCR runtime error
        return ""

    return "\n".join(ocr_text)

def extract_text_from_cv(file_obj) -> str:
    """
    Extracts text from uploaded CV files (.pdf, .docx, .txt, .csv) with:
    - File size and format validation
    - Scanned PDF detection
    - Graceful optional OCR fallback
    - Table content extraction for Word documents
    """
    validate_cv_file(file_obj)
    filename = getattr(file_obj, "name", "document").lower()
    raw_text = ""

    # Ensure stream is at start
    if hasattr(file_obj, "seek"):
        file_obj.seek(0)

    if filename.endswith(".pdf"):
        try:
            reader = PdfReader(file_obj)
            for page in reader.pages:
                t = page.extract_text()
                if t:
                    raw_text += t + "\n"
        except Exception as e:
            raise CVExtractionError(f"Failed to parse PDF document: {str(e)}")

        # Check if text is empty or too short (scanned PDF detection)
        alphanumeric_chars = len(re.sub(r'[^a-zA-Z0-9]', '', raw_text))
        if alphanumeric_chars < 50:
            # Scanned PDF detected, attempt optional OCR
            ocr_result = _attempt_ocr_on_pdf(reader)
            if ocr_result and len(re.sub(r'[^a-zA-Z0-9]', '', ocr_result)) >= 50:
                raw_text = ocr_result
            else:
                raise CVExtractionError(
                    "The uploaded PDF appears to be a scanned image with no selectable text. "
                    "Please upload a text-based PDF/DOCX or copy-paste your CV text directly."
                )

    elif filename.endswith(".docx") or filename.endswith(".doc"):
        if docx is None:
            raise CVExtractionError("Word document parsing library (python-docx) is not installed.")
        try:
            doc = docx.Document(file_obj)
            parts = []
            # Extract paragraphs
            for p in doc.paragraphs:
                if p.text.strip():
                    parts.append(p.text.strip())
            # Extract table cells (frequently used for skills/experience layouts)
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        parts.append(row_text)
            raw_text = "\n".join(parts)
        except Exception as e:
            raise CVExtractionError(f"Failed to parse Word document: {str(e)}")

    elif filename.endswith(".csv"):
        try:
            df = pd.read_csv(file_obj)
            raw_text = df.to_string()
        except Exception as e:
            raise CVExtractionError(f"Failed to parse CSV file: {str(e)}")

    else:
        # Plain text
        try:
            content = file_obj.read()
            if isinstance(content, str):
                raw_text = content
            else:
                raw_text = content.decode("utf-8", errors="replace")
        except Exception as e:
            raise CVExtractionError(f"Failed to read text file: {str(e)}")

    clean_res = raw_text.strip()
    if not clean_res:
        raise CVExtractionError("The uploaded file does not contain readable text.")

    return clean_res
